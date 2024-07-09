import glob
import logging as log
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from tqdm import tqdm

from .dqe_gt import MockDqeGT
from .dqe_handler import SWC
from .dqe_io import DqeInput, DqeModel, DqeOuput, DqeProcess
from .utils import read_ini

# Global
RK = "R"
WK = "W"
POS = "positive"
NEG = "negative"
PASS = "PASS"
FAIL = "FAIL"
IMG_EXTS = [".png", ".jpg"]
JSON_EXT = ".json"


@dataclass
class XmlBasicInfo:
    disk_name: str
    keyword: str
    total_num: int
    positive_num: int
    negative_num: int
    rate: float


@dataclass
class XmlSuccessData:
    file_name: str
    detected: str
    result: str


@dataclass
class XmlFailedData:
    file_name: str
    error: str


class DqeXmlHandler:
    XML_EXT: str = ".xlsx"

    def __init__(self, output_dir: str, output_name: Optional[str] = None) -> None:
        if output_name is None:
            now = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_name = f"{now}{self.XML_EXT}"
        self.xml_path = Path(output_dir) / output_name
        self.xml_path.parent.mkdir(parents=True, exist_ok=True)

        self.wb = Workbook()
        self.first_ws = self.wb.active

    def add_page(self, title: str, contents: List[list]):
        cur_ws = self.wb.create_sheet(title=title)
        for content in contents:
            cur_ws.append(content)
        self.adjust_worksheet(cur_ws=cur_ws)

    def adjust_worksheet(self, cur_ws, scale: float = 1.2, bias: int = 1):
        # 自動調整列寬
        for col in cur_ws.columns:
            max_length = 0
            column = col[0].column_letter  # Get the column name
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except BaseException:
                    pass
            adjusted_width = (max_length + bias) * scale
            cur_ws.column_dimensions[column].width = adjusted_width

    def save_xml(self):
        self.wb.save(self.xml_path)

    def print_xml(self):
        pass


def dc2list(dataclass_instance) -> list:
    return list(asdict(dataclass_instance).values())


class SWCForValidator(SWC):
    def __init__(self, config: Dict) -> None:
        self.config = config
        self.dinputs: Optional[List[DqeInput]] = None
        self.models: Optional[List[DqeModel]] = None

        self.GT = MockDqeGT(disk_name=config["test-disk"]["disk_name"])
        self.xml_handler = DqeXmlHandler(output_dir=config["output"]["output_dir"])

    def _get_inputs(self, config: dict) -> List[DqeInput]:
        dprocess = DqeProcess(
            module_name="process", module_path=config["process"]["module_path"]
        )
        # Capture all image with different extensions
        image_list = []
        for ext in IMG_EXTS:
            glob_path = str(Path(config["input"]["input_dir"]) / "**" / rf"*{ext}")
            cur_image_list = glob.glob(glob_path, recursive=True)
            if cur_image_list:
                log.debug(f"Find {len(cur_image_list)} '{ext}' images")
                image_list += cur_image_list

        # Wrap to DqeInput
        dinputs = []
        for image in image_list:
            try:
                dinputs.append(DqeInput(image_path=image, dqe_process=dprocess))
            except BaseException as e:
                log.warning(f"The input image is wrong: {image} ({e})")
                continue
        return dinputs

    def _get_models(self, config) -> Dict[str, DqeModel]:
        from collections import defaultdict

        models = defaultdict()
        keywords, config_keywords = [RK, WK], ["read", "write"]
        for model_key, config_key in zip(keywords, config_keywords):
            title = f"model.{config_key}"
            if title not in config:
                log.warning(f"not setup {config_key} in config")
                continue
            try:
                if int(config[title]["enable"]) == 0:
                    log.warning(f"disable model {model_key}")
                    continue

                models[model_key] = DqeModel(config[title])
            except BaseException:
                pass
                # log.warning("The model setting is wrong. please check again")
                # raise err
        log.info(f"Get {', '.join(models.keys())} Model")
        return models

    def inference(self):
        assert self.models, "Model not loaded, please use SWC.load()"
        assert len(self.models) == 1, "Only support 1 model in validator.ini"
        model_keyword = list(self.models.keys())[0]
        # Filter input
        log.info("Filter wrong input ...")
        wrong_inputs: List[list] = []
        need_remove = []
        for dinput in self.dinputs:
            if dinput.keyword != model_keyword:
                wrong_inputs.append(
                    XmlFailedData(file_name=dinput.path, error="wrong_key")
                )
                need_remove.append(dinput)
                continue
        for dinput in need_remove:
            self.dinputs.remove(dinput)
        log.warning(f"Found {len(wrong_inputs)} wrong keyword images")

        # Do inference
        log.info("Start inference")
        output_data: List[list] = []
        for dinput in tqdm(self.dinputs):
            # Inference with callback
            self.models[dinput.keyword].inference(dinput)
            doutput: DqeOuput = self.models[dinput.keyword].output
            label, stats = None, NEG
            if doutput.output:
                index, label, conf = doutput.output[0]  # Top 1
                stats = POS if self.GT.compare(label) else NEG
                output_data.append(
                    XmlSuccessData(file_name=dinput.path, detected=label, result=stats)
                )

        # For XML
        num_tot = len(output_data)
        num_pos = len([data for data in output_data if data.result == POS])
        num_neg = num_tot - num_pos
        rate = int((num_pos / num_tot) * 100) if num_tot and num_pos else 0
        basic_info = XmlBasicInfo(
            disk_name=self.GT.ans,
            keyword=model_keyword,
            total_num=num_tot,
            positive_num=num_pos,
            negative_num=num_neg,
            rate=rate,
        )
        self.xml_handler.first_ws.title = "Overview"
        self.xml_handler.first_ws.append(["Category", "Content"])
        for key, val in zip(
            ["Disk", "Mode", "Total", "Positive", "Negative", "Rate"],
            dc2list(basic_info),
        ):
            self.xml_handler.first_ws.append([key, val])
        self.xml_handler.adjust_worksheet(cur_ws=self.xml_handler.first_ws, bias=4)

        chart = PieChart()
        chart.title = "Postive / Negative"
        categories = Reference(
            self.xml_handler.first_ws, min_col=1, min_row=5, max_row=6
        )
        data = Reference(self.xml_handler.first_ws, min_col=2, min_row=5, max_row=6)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(categories)
        self.xml_handler.first_ws.add_chart(chart, "E5")

        self.xml_handler.add_page(
            title="Results",
            contents=[["File Path", "Detected", "Result"]]
            + [dc2list(dc) for dc in output_data],
        )
        self.xml_handler.add_page(
            title="Failed",
            contents=[["File Path", "Error Message"]]
            + [dc2list(dc) for dc in wrong_inputs],
        )

        self.xml_handler.save_xml()


if __name__ == "__main__":
    config = read_ini(config_file="validator.ini")
    swcv = SWCForValidator(config=config)
    swcv.load()
    swcv.inference()
